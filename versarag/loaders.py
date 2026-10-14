"""Document loaders keyed by file extension.

Register your own with ``@register_loader(".ext")``.
"""

from __future__ import annotations

import csv
import json
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Iterable, Iterator

from .types import Document

Loader = Callable[[Path], Iterable[Document]]
_LOADERS: dict[str, Loader] = {}

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".tex", ".org",
    ".py", ".r", ".rmd", ".qmd", ".sql", ".js", ".ts", ".java", ".go", ".rs",
    ".c", ".cpp", ".h", ".sh", ".yaml", ".yml", ".toml", ".ini", ".cfg",
}


def register_loader(*extensions: str) -> Callable[[Loader], Loader]:
    def deco(fn: Loader) -> Loader:
        for ext in extensions:
            _LOADERS[ext.lower()] = fn
        return fn

    return deco


def supported_extensions() -> list[str]:
    return sorted(_LOADERS)


def _meta(path: Path, **extra: object) -> dict[str, object]:
    return {"source": str(path), "filename": path.name, **extra}


@register_loader(*TEXT_EXTENSIONS)
def load_text(path: Path) -> Iterator[Document]:
    yield Document(path.read_text(encoding="utf-8", errors="replace"), _meta(path))


@register_loader(".pdf")
def load_pdf(path: Path) -> Iterator[Document]:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ImportError("PDF support needs `pip install pypdf` (or versarag[docs])") from e
    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            yield Document(text, _meta(path, page=i))


@register_loader(".docx")
def load_docx(path: Path) -> Iterator[Document]:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise ImportError("DOCX support needs `pip install python-docx` (or versarag[docs])") from e
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    yield Document("\n\n".join(parts), _meta(path))


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: object) -> None:
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


@register_loader(".html", ".htm")
def load_html(path: Path) -> Iterator[Document]:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    lines = [" ".join(line.split()) for line in "".join(parser.parts).splitlines()]
    yield Document("\n".join(line for line in lines if line), _meta(path))


@register_loader(".csv", ".tsv")
def load_csv(path: Path) -> Iterator[Document]:
    """Each row becomes one document of ``column: value`` lines."""
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        for i, row in enumerate(csv.DictReader(f, delimiter=delimiter), start=1):
            text = "\n".join(f"{k}: {v}" for k, v in row.items() if v not in (None, ""))
            if text:
                yield Document(text, _meta(path, row=i))


@register_loader(".json", ".jsonl")
def load_json(path: Path) -> Iterator[Document]:
    """JSONL: one document per line. JSON: a list -> one doc per item, else one doc.

    If items are objects with a ``text`` (or ``content``) field, it is used as the
    text and the remaining scalar fields become metadata.
    """
    raw = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".jsonl":
        items = [json.loads(line) for line in raw.splitlines() if line.strip()]
    else:
        data = json.loads(raw)
        items = data if isinstance(data, list) else [data]
    for i, item in enumerate(items, start=1):
        if isinstance(item, dict) and ("text" in item or "content" in item):
            text = str(item.get("text", item.get("content")))
            extra = {
                k: v for k, v in item.items()
                if k not in ("text", "content") and isinstance(v, (str, int, float, bool))
            }
            yield Document(text, {**_meta(path, item=i), **extra})
        else:
            yield Document(json.dumps(item, ensure_ascii=False, indent=1), _meta(path, item=i))


def iter_files(paths: Iterable[str | Path], recursive: bool = True) -> Iterator[Path]:
    for p in map(Path, paths):
        if p.is_dir():
            files = p.rglob("*") if recursive else p.glob("*")
            for f in sorted(files):
                if f.is_file() and f.suffix.lower() in _LOADERS and not any(
                    part.startswith(".") for part in f.relative_to(p).parts
                ):
                    yield f
        elif p.is_file():
            yield p
        else:
            raise FileNotFoundError(p)


def load_file(path: str | Path) -> list[Document]:
    path = Path(path)
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(
            f"No loader for {path.suffix!r} ({path}). Supported: {supported_extensions()}"
        )
    return list(loader(path))
